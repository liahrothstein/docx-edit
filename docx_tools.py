# docx_tools.py
import sys
from docx import Document

def outline(path):
    """Структура: заголовки и их индексы."""
    doc = Document(path)
    for i, p in enumerate(doc.paragraphs):
        n = p.style.name.lower()
        if "heading" in n or "заголовок" in n:
            print(f"[{i}] {p.style.name}: {p.text}")

def find(path, query):
    doc = Document(path)
    for i, p in enumerate(doc.paragraphs):
        if query.lower() in p.text.lower():
            print(f"[{i}] {p.text[:150]}")

def show(path, index):
    doc = Document(path)
    for r in doc.paragraphs[int(index)].runs:
        fmt = " ".join(f for f, v in
            [("bold", r.bold), ("italic", r.italic),
             ("sub", r.font.subscript)] if v)
        print(repr(r.text), fmt)

def _replace_in_paragraph(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    if old in p.text and p.runs:
        full = p.text.replace(old, new)
        for r in p.runs:
            r.text = ""
        p.runs[0].text = full
        return True
    return False

def replace(path, old, new, out=None):
    doc = Document(path)
    targets = list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                targets += cell.paragraphs
    for s in doc.sections:
        targets += s.header.paragraphs + s.footer.paragraphs
    n = sum(1 for p in targets if _replace_in_paragraph(p, old, new))
    doc.save(out or path)   # по умолчанию — в исходник, бэкапом служит git
    print(f"Заменено абзацев: {n} -> {out or path}")

def text(path, start, end=None):
    """Текст абзацев в диапазоне индексов [start, end)."""
    doc = Document(path)
    start, end = int(start), int(end) if end else int(start) + 1
    for i, p in enumerate(doc.paragraphs[start:end], start=start):
        if p.text.strip():
            print(f"[{i}] {p.text}")

def settext(path, index, new_text):
    """Перезаписать абзац целиком (форматирование первого run сохраняется)."""
    doc = Document(path)
    p = doc.paragraphs[int(index)]
    for r in p.runs:
        r.text = ""
    if p.runs: p.runs[0].text = new_text
    else: p.add_run(new_text)
    doc.save(path)
    print(f"Абзац [{index}] переписан -> {path}")

def setbody(path, index, new_text):
    """Перезаписать тело абзаца, сохранив первый run (жирное «Ответ:»)."""
    doc = Document(path)
    p = doc.paragraphs[int(index)]
    if not p.runs or "Ответ" not in p.runs[0].text:
        print(f"[{index}] первый run не содержит 'Ответ' — пропуск, проверить show")
        return
    for r in p.runs[1:]:
        r.text = ""
    sep = "" if p.runs[0].text.endswith((" ", "\u00A0")) else " "
    if len(p.runs) > 1:
        p.runs[1].text = sep + new_text
    else:
        p.add_run(sep + new_text)
    doc.save(path)
    print(f"Тело абзаца [{index}] заменено -> {path}")

if __name__ == "__main__":
    cmd, path = sys.argv[1], sys.argv[2]
    if cmd == "outline": outline(path)
    elif cmd == "find":  find(path, sys.argv[3])
    elif cmd == "show":  show(path, sys.argv[3])
    elif cmd == "text":    text(path, sys.argv[3], sys.argv[4] if len(sys.argv) > 4 else None)
    elif cmd == "settext": settext(path, sys.argv[3], sys.argv[4])
    elif cmd == "setbody": setbody(path, sys.argv[3], sys.argv[4])
    elif cmd == "replace":
        replace(path, sys.argv[3], sys.argv[4],
                sys.argv[5] if len(sys.argv) > 5 else None)